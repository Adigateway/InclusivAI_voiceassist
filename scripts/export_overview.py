"""Export data/raw/services_raw.json to a human-readable overview (data/exports/)."""
import json, csv
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

ROOT = Path(__file__).resolve().parent.parent
d = json.load(open(ROOT / "data/raw/services_raw.json", encoding="utf8"))
hdr = ["section", "slug", "title", "url", "chars", "phone_count", "text"]
rows = [[s["section"], s["slug"], s["title"], s["url"], len(s["text"]),
         sum(l["url"].startswith("tel:") for l in s["external_links"]), s["text"]] for s in d]
lhdr = ["section", "slug", "type", "text", "url"]
links = [[s["section"], s["slug"], l["url"].split(":")[0] if l["url"].startswith(("tel:", "mailto:")) else "web",
          l["text"], l["url"]] for s in d for l in s["external_links"]]

with open(ROOT / "data/exports/services_pages.csv", "w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f); w.writerow(hdr); w.writerows(rows)

wb = Workbook(); wb.active.title = "Pages"
for name, h, r, widths in [("Pages", hdr, rows, [22, 28, 40, 50, 8, 12, 120]), ("Contacts", lhdr, links, [22, 28, 8, 50, 70])]:
    ws = wb[name] if name in wb.sheetnames else wb.create_sheet(name)
    ws.append(h); [ws.append(x) for x in r]
    for c in ws[1]: c.font = Font(bold=True)
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    for i, wd in enumerate(widths): ws.column_dimensions[chr(65 + i)].width = wd
    if name == "Pages":
        for c in ws["G"][1:]: c.alignment = Alignment(wrap_text=True, vertical="top")
wb.save(ROOT / "data/exports/services_overview.xlsx")
print(len(rows), "pages,", len(links), "links")
